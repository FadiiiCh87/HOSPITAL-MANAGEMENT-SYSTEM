from app.services.report_service import (
    get_revenue_report,
    get_appointment_report,
    get_all_appointments
)
from app.utils.permissions import admin_required
from flask import (
    request,
    render_template,
    Blueprint,
    send_file
)
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph
)
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors

from reportlab.lib.styles import getSampleStyleSheet

from reportlab.lib.units import inch
import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from flask_login import login_required



report_bp = Blueprint(
    "report",
    __name__
)

@report_bp.route("/reports/revenue")
@login_required
@admin_required
def revenue_report():

    from_date = request.args.get("from_date")
    to_date = request.args.get("to_date")

    report = get_revenue_report(
        from_date,
        to_date
    )

    return render_template(
        "reports/revenue.html",
        from_date=from_date,
        to_date=to_date,
        **report
    )
# ==========================
# Appointment Report
# ==========================
@report_bp.route("/reports/appointments")
@login_required
@admin_required
def appointment_report():

    from_date = request.args.get("from_date")
    to_date = request.args.get("to_date")

    report = get_appointment_report(
        from_date,
        to_date
    )

    return render_template(
        "reports/appointments.html",
        from_date=from_date,
        to_date=to_date,
        **report
    )
# ==========================
# Export Appointment Report (Excel)
# ==========================
@report_bp.route("/reports/appointments/export")
@login_required
@admin_required
def export_appointments_excel():

    wb = Workbook()
    ws = wb.active
    ws.title = "Appointments"

    # ==========================
    # Report Title
    # ==========================

    ws.merge_cells("A1:F1")
    ws["A1"] = "Hospital Management System"
    ws["A1"].font = Font(bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:F2")
    ws["A2"] = "Appointment Report"
    ws["A2"].font = Font(bold=True, size=14)
    ws["A2"].alignment = Alignment(horizontal="center")

    # ==========================
    # Table Headers
    # ==========================

    headers = [
        "ID",
        "Patient",
        "Doctor",
        "Date",
        "Time",
        "Status"
    ]

    for col, header in enumerate(headers, start=1):

        cell = ws.cell(row=3, column=col)
        cell.value = header

        cell.fill = PatternFill(
            fill_type="solid",
            start_color="1F4E78",
            end_color="1F4E78"
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.alignment = Alignment(horizontal="center")

    # ==========================
    # Appointment Data
    # ==========================

    appointments = get_all_appointments()

    row = 4

    for appointment in appointments:

        ws.cell(row=row, column=1).value = appointment.id
        ws.cell(row=row, column=2).value = appointment.patient.full_name
        ws.cell(row=row, column=3).value = appointment.doctor.full_name
        ws.cell(row=row, column=4).value = appointment.appointment_date.strftime("%d-%m-%Y")
        ws.cell(row=row, column=5).value = appointment.appointment_time.strftime("%I:%M %p")
        ws.cell(row=row, column=6).value = appointment.status

        row += 1

    # ==========================
    # Auto Width
    # ==========================

    for col in range(1, 7):

        column_letter = get_column_letter(col)
        max_length = 0

        for cell in ws[column_letter]:

            if cell.value:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[column_letter].width = max_length + 5

    # ==========================
    # Freeze Header
    # ==========================

    ws.freeze_panes = "A4"

    # ==========================
    # Auto Filter
    # ==========================

    ws.auto_filter.ref = f"A3:F{ws.max_row}"

    # ==========================
    # Download
    # ==========================

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Appointment_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
# ==========================
# Export Appointment Report (PDF)
# ==========================
@report_bp.route("/reports/appointments/pdf")
@login_required
@admin_required
def export_appointments_pdf():

    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4
    )

    styles = getSampleStyleSheet()

    elements = []

    title = Paragraph(
        "<b>Hospital Management System</b>",
        styles["Title"]
    )

    subtitle = Paragraph(
        "<b>Appointment Report</b>",
        styles["Heading2"]
    )

    elements.append(title)
    elements.append(subtitle)

    appointments = get_all_appointments()
    # ==========================
    # Table Data
    # ==========================

    data = [

    [
        "ID",
        "Patient",
        "Doctor",
        "Date",
        "Time",
        "Status"
        ]

        ]

    for appointment in appointments:

        data.append([

        str(appointment.id),

        appointment.patient.full_name,

        appointment.doctor.full_name,

        appointment.appointment_date.strftime("%d-%m-%Y"),

        appointment.appointment_time.strftime("%I:%M %p"),

        appointment.status

        ])
        table = Table(data)
        table.setStyle(TableStyle([

        ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),

        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),

        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),

        ("ALIGN", (0, 0), (-1, -1), "CENTER"),

        ("GRID", (0, 0), (-1, -1), 1, colors.black),

        ("BACKGROUND", (0, 1), (-1, -1), colors.beige),

        ("BOTTOMPADDING", (0, 0), (-1, 0), 10)

        ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)

    return send_file(
    buffer,
    as_attachment=True,
    download_name="Appointment_Report.pdf",
    mimetype="application/pdf"

)
    